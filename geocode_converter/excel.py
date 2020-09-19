import openpyxl


def load_file(path: str) -> list:
    wb = openpyxl.open(path)
    ws = wb.active

    data_list = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if None not in row:
            data_list.append(
                {
                    "id": row[0],
                    "latitude": float(str(row[1]).replace(",", ".")),
                    "longitude": float(str(row[2]).replace(",", ".")),
                }
            )

    return data_list


def output_data(data: list, file_path: str):
    labels = list(data[0].keys())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OUTPUT"

    ws.append(labels)
    for item in data:
        values = (item[key] for key in labels)
        ws.append(values)

    tbl = openpyxl.worksheet.table.Table(
        displayName="output", ref=ws.dimensions
    )

    ws.add_table(tbl)
    wb.save(file_path)